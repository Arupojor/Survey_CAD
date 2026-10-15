"""Survey import and exact piecewise-linear earthwork calculations. No Excel macros run."""
from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from bisect import bisect_right, bisect_left
import csv, math, re
import openpyxl

EPS = 1e-9

class InputError(ValueError):
    pass

def numeric(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)

def number(v, label):
    if not numeric(v):
        raise InputError(f'{label}: expected a finite number, found {v!r}. Recalculate and save formula cells in Excel before importing.')
    return float(v)

def clean_points(points, label):
    if len(points) < 2:
        raise InputError(f'{label}: at least two points are required.')
    result = [(number(x, label+' offset'), number(y, label+' level')) for x,y in points]
    if any(b[0] < a[0] for a,b in zip(result,result[1:])) or result[-1][0]-result[0][0]<=EPS:
        raise InputError(f'{label}: offsets must not decrease and must span a positive width. Points are not silently reordered.')
    return result

def interp(points, x, side='right'):
    if x < points[0][0]-EPS or x > points[-1][0]+EPS:
        raise InputError(f'Offset {x:g} is outside the surveyed profile.')
    xs=[p[0] for p in points]
    left=bisect_left(xs,x); right=bisect_right(xs,x)
    if left<right: return points[left if side=='left' else right-1][1]
    i = max(0, min(len(points)-2, right-1))
    a,b = points[i:i+2]
    return a[1]+(b[1]-a[1])*(x-a[0])/(b[0]-a[0])

def poly_area(points):
    if len(points)<3:return 0.
    ox,oy=points[0]; p=[(x-ox,y-oy) for x,y in points]
    return abs(math.fsum(a[0]*b[1]-b[0]*a[1] for a,b in zip(p,p[1:]+p[:1])))/2

def integral(points, datum=0):
    return sum((b[0]-a[0])*(a[1]+b[1]-2*datum)/2 for a,b in zip(points,points[1:]))

@dataclass
class Section:
    chainage: float
    pre: list
    post: list
    source: str = ''
    label: str = ''
    warnings: list = field(default_factory=list)
    fill_polygons: list = field(default_factory=list)
    cut_polygons: list = field(default_factory=list)
    fill: float = 0
    cut: float = 0
    net: float = 0
    pre_integral: float = 0
    post_integral: float = 0
    cached_net: float | None = None
    mean_fill: float = 0
    mean_cut: float = 0
    mean_net: float = 0

    def calculate(self):
        self.warnings=[]
        self.pre = clean_points(self.pre, f'Ch {self.chainage:g} pre-work')
        self.post = clean_points(self.post, f'Ch {self.chainage:g} post-work')
        for name,pts in [('pre',self.pre),('post',self.post)]:
            for a,b in zip(pts,pts[1:]):
                if a[0]==b[0]:self.warnings.append(f'{name}: repeated offset {a[0]:g} m retained as a vertical segment ({a[1]:g} to {b[1]:g} RL). Check source entry; no correction assumed.')
        if abs(self.pre[0][0]-self.post[0][0])>EPS or abs(self.pre[-1][0]-self.post[-1][0])>EPS:
            raise InputError(f'Ch {self.chainage:g}: pre/post profiles must have the same first and last offsets. No extrapolation or silent trimming is performed.')
        xs = sorted(set(x for x,_ in self.pre+self.post))
        crossings=[]
        for a,b in zip(xs,xs[1:]):
            da=interp(self.post,a)-interp(self.pre,a)
            db=interp(self.post,b,'left')-interp(self.pre,b,'left')
            if da*db<0:
                crossings.append(a+(b-a)*da/(da-db))
        xs=sorted(set(xs+crossings))
        intervals=[((a,interp(self.pre,a),interp(self.post,a)),(b,interp(self.pre,b,'left'),interp(self.post,b,'left'))) for a,b in zip(xs,xs[1:])]
        self.fill_polygons=[]; self.cut_polygons=[]
        # Group adjacent strips into simple polygons, splitting at every zero crossing.
        run=[]; sign=0
        def finish():
            if len(run)>1:
                pts=[(x,post) for x,pre,post in run]+[(x,pre) for x,pre,post in reversed(run)]
                dedup=[]
                for p in pts:
                    if not dedup or math.dist(p,dedup[-1])>EPS: dedup.append(p)
                if len(dedup)>1 and math.dist(dedup[0],dedup[-1])<EPS: dedup.pop()
                if len(dedup)>=3 and poly_area(dedup)>1e-12:
                    (self.fill_polygons if sign>0 else self.cut_polygons).append(dedup)
        fill_integral=cut_integral=0.
        for a,b in intervals:
            depth=(a[2]-a[1]+b[2]-b[1])/2
            area=depth*(b[0]-a[0]); s=1 if depth>EPS else -1 if depth<-EPS else 0
            fill_integral+=max(area,0); cut_integral+=max(-area,0)
            if s!=sign or abs(a[2]-a[1])<EPS:
                finish(); run=[]; sign=s
            if s:
                if not run: run=[a]
                elif run[-1]!=a: run.append(a)
                run.append(b)
        finish()
        self.fill=sum(poly_area(p) for p in self.fill_polygons)
        self.cut=sum(poly_area(p) for p in self.cut_polygons)
        self.net=self.fill-self.cut
        self.mean_fill=fill_integral;self.mean_cut=cut_integral;self.mean_net=fill_integral-cut_integral
        self.pre_integral=integral(self.pre); self.post_integral=integral(self.post)
        if max(abs(self.fill-fill_integral),abs(self.cut-cut_integral),abs(self.net-(self.post_integral-self.pre_integral)))>1e-6:
            raise InputError(f'Internal area check failed at Ch {self.chainage:g}.')
        if any(abs(interp(self.pre,x)-interp(self.post,x))>1e-6 for x in [xs[0],xs[-1]]):
            self.warnings.append('Profiles do not meet at a survey edge; fill/cut is bounded by that surveyed vertical edge.')
        if self.cached_net is not None and min(abs(self.net-self.cached_net),abs(-self.net-self.cached_net))>1e-5:
            self.warnings.append(f'Workbook cached area {self.cached_net:.6f} m2; recomputed signed net {self.net:.6f} m2 (fill minus cut). Magnitude discrepancy {abs(self.net)-abs(self.cached_net):+.6f} m2. Check workbook arithmetic.')
        return self

@dataclass
class Survey:
    sections: list
    source: str
    mode: str
    warnings: list = field(default_factory=list)
    reference_volume: float | None = None
    design_reaches: list = field(default_factory=list)
    design_basis: str = ''

    def validate(self):
        if not self.sections: raise InputError('No complete sections were found in the selected source.')
        self.sections.sort(key=lambda s:s.chainage)
        for a,b in zip(self.sections,self.sections[1:]):
            if b.chainage-a.chainage<EPS: raise InputError(f'Duplicate chainage {a.chainage:g} m. Use one alignment/bank per import.')
        for s in self.sections: s.calculate()
        return self

def cell(ws, row, col):
    c=ws.cell(row,col)
    return number(c.value,f"'{ws.title}'!{c.coordinate}")

def read_area(wb, source, sheet=None):
    name=sheet or next((s for s in wb.sheetnames if s.strip().lower() in ('area','cross section')),None)
    if not name: raise InputError('No Area worksheet found. Choose the flat-table mode for a different Excel layout.')
    ws=wb[name]; rows=list(ws.values); sections=[]; current=None
    for ri,row in enumerate(rows,1):
        if len(row)<9: continue
        if numeric(row[0]) and numeric(row[1]):
            current=Section(float(row[1])*1000,[],[],f'{name}!row {ri}',str(row[0])); sections.append(current)
        if current is None: continue
        if isinstance(row[2],str) and 'total' in row[2].lower():
            if len(row)>12 and numeric(row[12]): current.cached_net=float(row[12])
            current=None; continue
        for xc,yc,key in [(2,3,'pre'),(7,8,'post')]:
            x,y=row[xc],row[yc]
            if x is None and y is None: continue
            if isinstance(x,str) and not x.startswith('#'): continue
            getattr(current,key).append((number(x,f'{name} row {ri} offset'),number(y,f'{name} row {ri} level')))
    reference=None
    if 'vol' in wb:
        for row in wb['vol'].values:
            if isinstance(row[0],str) and 'total' in row[0].lower() and len(row)>5 and numeric(row[5]): reference=float(row[5])
    result=Survey(sections,source,'Pre/post coordinates - Area sheet',[
        'Area-sheet chainages are interpreted as kilometres and converted to metres.',
        'Formula input cells use their last saved Excel values. Recalculate and save Excel before import; this app does not run Excel formulas.',
        'Post-work coordinates are used as provided. The workbook may describe design rather than measured as-built levels.'
    ],reference)
    if 'Gradation (2)' in wb:
        grade=wb['Gradation (2)']
        for xc,yc in [(13,14),(16,17),(19,20),(22,23)]:
            pairs=[(float(row[xc-1])*1000,float(row[yc-1])) for row in grade.iter_rows(min_row=3,values_only=True) if len(row)>=yc and numeric(row[xc-1]) and numeric(row[yc-1])]
            if len(pairs)>1:
                a,y=pairs[0];b,z=pairs[-1]
                if b>a and all(abs(rl-(y+(z-y)*(ch-a)/(b-a)))<1e-7 for ch,rl in pairs):
                    result.design_reaches.append((round(a,6),round(b,6),y,z))
        if result.design_reaches:
            result.design_basis='Gradation (2): M:N, P:Q, S:T, V:W; chainage km converted to m. Reach resets retained; review continuity.'
    return result

def read_design(wb, source):
    if 'XS Input' not in wb or 'XS Areas' not in wb: raise InputError('Design mode requires XS Input and XS Areas worksheets.')
    ws=wb['XS Input']; formation=cell(ws,5,2); width=cell(ws,6,2); centre=cell(ws,7,2)
    country=cell(ws,8,2); river=cell(ws,9,2)
    side=str(ws.cell(10,2).value).strip().lower()
    if not (side.startswith('low') or side.startswith('high')): raise InputError('XS Input!B10 must be Low offset or High offset.')
    overrides={}
    for ri,row in enumerate(wb['XS Areas'].iter_rows(min_row=6,values_only=True),6):
        if numeric(row[0]) and numeric(row[1]):
            key=round(float(row[1])*1000,8)
            if key in overrides: raise InputError(f'Duplicate design override chainage {key:g}.')
            overrides[key]=[default if v is None or v=='' else number(v,f'XS Areas row {ri} override') for v,default in zip(row[5:10],[centre,formation,width,country,river])]
    grouped={}
    for ri,row in enumerate(ws.iter_rows(min_row=12,values_only=True),12):
        a,x,y=row[:3]
        if all(v is None or v=='' for v in (a,x,y)): continue
        ch=number(a,f'XS Input!A{ri}')*1000
        grouped.setdefault(round(ch,8),[]).append((number(x,f'XS Input!B{ri}'),number(y,f'XS Input!C{ri}')))
    sections=[]
    for ch,pts in grouped.items():
        pts=clean_points(pts,f'Ch {ch:g} ground')
        c,f,w,cs,rs=overrides.get(ch,[centre,formation,width,country,river])
        if min(w,cs,rs)<=0: raise InputError(f'Ch {ch:g}: crest width and slopes must be positive.')
        left,right=c-w/2,c+w/2; sl,sr=(cs,rs) if side.startswith('low') else (rs,cs)
        def design(x): return f+(x-left)/sl if x<left else f-(x-right)/sr if x>right else f
        if left<pts[0][0] or right>pts[-1][0] or design(pts[0][0])>pts[0][1]+EPS or design(pts[-1][0])>pts[-1][1]+EPS:
            raise InputError(f'Ch {ch:g}: survey does not extend beyond the design toes/crest. Add ground points before calculating a complete section.')
        xs=sorted(set([x for x,_ in pts]+[left,right])); crossings=[]
        for a,b in zip(xs,xs[1:]):
            da=design(a)-interp(pts,a); db=design(b)-interp(pts,b)
            if da*db<0: crossings.append(a+(b-a)*da/(da-db))
        xs=sorted(set(xs+crossings)); post=[(x,max(interp(pts,x),design(x))) for x in xs]
        sections.append(Section(ch,pts,post,'XS Input + XS Areas',str(len(sections)+1)))
    ref=wb['XS Volume']['B4'].value if 'XS Volume' in wb else None
    return Survey(sections,source,'Design envelope - XS Input / XS Areas',[
        'Design mode reconstructs a fill-only post profile from ground, crest level/width, centre offset and side slopes.',
        'Ground above the design is retained. This is a design quantity, not measured post-work earthwork.',
        'Chainages are converted from km to m. Any formula inputs use saved Excel values.'
    ],float(ref) if numeric(ref) else None)

def read_flat(rows,source,sheet):
    # Explicit long-form contract keeps separate pre/post offsets and never guesses units.
    rows=list(rows); header=None; idx=None
    normalize=lambda x: re.sub(r'[^a-z0-9]','',str(x).lower())
    for i,row in enumerate(rows[:30]):
        hs=[normalize(v) for v in row]
        if 'profile' in hs and 'offsetm' in hs and 'rlm' in hs and ('chainagem' in hs or 'chainagekm' in hs): header=hs; idx=i; break
    if header is None: raise InputError('Flat table needs headers Chainage_m (or Chainage_km), Profile, Offset_m, RL_m. Profile must be Pre or Post.')
    ci=header.index('chainagem') if 'chainagem' in header else header.index('chainagekm'); factor=1 if 'chainagem' in header else 1000
    pi=header.index('profile'); xi=header.index('offsetm'); yi=header.index('rlm'); groups={}
    for ri,row in enumerate(rows[idx+1:],idx+2):
        if all(v is None or str(v).strip()=='' for v in row): continue
        if len(row)<=max(ci,pi,xi,yi): raise InputError(f'Row {ri} is incomplete.')
        def val(v,label):
            try: v=float(v)
            except (TypeError,ValueError): raise InputError(f'Row {ri}: {label} must be numeric.')
            return number(v,f'row {ri} {label}')
        ch=val(row[ci],'chainage')*factor; profile=str(row[pi]).strip().lower()
        if profile not in ('pre','post'): raise InputError(f'Row {ri}: Profile must be Pre or Post.')
        key=round(ch,8); s=groups.setdefault(key,Section(key,[],[],sheet or 'CSV',str(len(groups)+1)))
        getattr(s,profile).append((val(row[xi],'offset'),val(row[yi],'RL')))
    return Survey(list(groups.values()),source,'Flat pre/post table',[
        'Offsets and levels are metres; chainage units are taken from the explicit column header.',
        'Use one alignment/bank per import. Flat-table formula inputs require saved Excel values.'
    ])

def load_survey(path,mode='Area',sheet=None):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f: survey=read_flat(csv.reader(f),path.name,None)
    elif path.suffix.lower() in ('.xlsx','.xlsm'):
        wb=openpyxl.load_workbook(path,data_only=True,read_only=True,keep_links=False)
        try:
            # Missing cached formula pairs must not disappear as apparently blank points.
            formulas=openpyxl.load_workbook(path,data_only=False,read_only=True,keep_links=False)
            try:
                if mode=='Area':
                    target=sheet or next((n for n in wb.sheetnames if n.strip().lower() in ('area','cross section')),None)
                    specs=[(target,{1,2,3,4,8,9})] if target else []
                elif mode=='Design': specs=[('XS Input',{1,2,3}),('XS Areas',{2,6,7,8,9,10})]
                else: specs=[(sheet or wb.sheetnames[0],None)]
                for name,cols in specs:
                    if name not in wb:continue
                    for rawrow,cachedrow in zip(formulas[name].iter_rows(),wb[name].iter_rows()):
                        for column,(raw,cached) in enumerate(zip(rawrow,cachedrow),1):
                            if raw.data_type=='f' and (cols is None or column in cols) and cached.value is None and cached.data_type not in ('s','str','inlineStr'):
                                raise InputError(f"'{name}'!{raw.coordinate}: formula has no saved value. Recalculate and save this workbook in Excel before importing.")
            finally:formulas.close()
            if mode=='Area': survey=read_area(wb,path.name,sheet)
            elif mode=='Design': survey=read_design(wb,path.name)
            else:
                selected=sheet or wb.sheetnames[0]
                survey=read_flat(wb[selected].values,path.name,selected)
        finally: wb.close()
    else: raise InputError('Choose an .xlsx, .xlsm or .csv file. Legacy .xls files must first be saved as .xlsx.')
    return survey.validate()

def exclusions(text):
    result=[]
    for part in re.split('[;\n]+',text.strip()):
        if not part.strip(): continue
        match=re.fullmatch(r'\s*(-?\d+(?:\.\d+)?)\s*:\s*(-?\d+(?:\.\d+)?)\s*',part)
        if not match: raise InputError('Excluded intervals must be start:end in metres, separated by semicolons. Example: 1200:1300; 2000:2050')
        a,b=map(float,match.groups())
        if b<=a: raise InputError('Each excluded interval must end after it starts.')
        result.append((a,b))
    result.sort(); merged=[]
    for a,b in result:
        if merged and a<=merged[-1][1]: merged[-1]=(merged[-1][0],max(b,merged[-1][1]))
        else: merged.append((a,b))
    return merged

def calculate_volumes(survey,start=None,end=None,excluded=(),extend=False):
    ss=survey.sections
    start=ss[0].chainage if start is None else number(start,'Start chainage')
    end=ss[-1].chainage if end is None else number(end,'End chainage')
    if end<=start: raise InputError('The end chainage must exceed the start chainage. At least two chainages, or an explicit extension, are needed for volume.')
    if not extend and (start<ss[0].chainage-EPS or end>ss[-1].chainage+EPS):
        raise InputError('Requested chainages extend beyond the survey. Enable the nearest-section extension assumption, or supply more sections.')
    cuts=sorted(set([start,end]+[s.chainage for s in ss if start<s.chainage<end]+[x for pair in excluded for x in pair if start<x<end]))
    def at(ch):
        if ch<=ss[0].chainage: return ss[0].fill,ss[0].cut
        if ch>=ss[-1].chainage: return ss[-1].fill,ss[-1].cut
        i=bisect_right([s.chainage for s in ss],ch)-1; a,b=ss[i:i+2]; t=(ch-a.chainage)/(b.chainage-a.chainage)
        return a.fill+(b.fill-a.fill)*t,a.cut+(b.cut-a.cut)*t
    rows=[]; cumulative=0.
    for a,b in zip(cuts,cuts[1:]):
        f1,c1=at(a); f2,c2=at(b); length=b-a
        skip=any(x-EPS<=(a+b)/2<=y+EPS for x,y in excluded)
        fv=0 if skip else (f1+f2)*length/2; cv=0 if skip else (c1+c2)*length/2
        cumulative+=fv-cv
        rows.append(dict(start=a,end=b,length=length,fill_start=f1,fill_end=f2,cut_start=c1,cut_end=c2,mean_fill=(f1+f2)/2,mean_cut=(c1+c2)/2,fill_volume=fv,cut_volume=cv,net_volume=fv-cv,cumulative_net=cumulative,excluded=skip,extended=a<ss[0].chainage-EPS or b>ss[-1].chainage+EPS))
    return dict(start=start,end=end,length=sum(r['length'] for r in rows if not r['excluded']),fill=sum(r['fill_volume'] for r in rows),cut=sum(r['cut_volume'] for r in rows),net=cumulative,rows=rows,excluded=list(excluded),extend=extend)
